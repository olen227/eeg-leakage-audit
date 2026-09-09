"""Сборка .docx из разметки Markdown под требования к оформлению ВКР.

ЗАЧЕМ ОТДЕЛЬНЫЙ СКРИПТ. Системная утилита `textutil` на этом тексте теряет таблицы и
заголовки (проверено: из 12 таблиц и 51 заголовка не сохранилось ни одного), поэтому
преобразование выполняется напрямую через python-docx с явным управлением стилями.

ЧТО ДЕЛАЕТ:
  * заголовки `#`/`##`/`###` — стилями Heading 1..3;
  * абзацы — основной текст, Times New Roman 14 пт, полуторный интервал,
    отступ первой строки 1,25 см, выравнивание по ширине;
  * строки вида `[РИСУНОК ...]` — отдельным абзацем по центру, курсивом,
    как место вставки иллюстрации;
  * таблицы Markdown — таблицами Word со стилем сетки;
  * цитаты `>` — с отступом и уменьшенным кеглем (служебные примечания);
  * блоки кода в тройных апострофах — ДОСЛОВНО, моноширинным шрифтом, по строке
    на абзац;
  * формульные вставки `$...$` разворачиваются в обычный текст: запись `0{,}1531`
    приводится к виду `0,1531`, служебные символы снимаются.

ПОЧЕМУ БЛОКИ КОДА ВЫДЕЛЕНЫ ОТДЕЛЬНО. Снятие формульной разметки убирает фигурные
скобки и обратную косую черту — в тексте это служебные знаки LaTeX, а в коде это
множества, словари, подстановки в f-строках и экранирование. Без отдельной ветви
листинги приложения превращались в синтаксически неверный код, склеенный в один
абзац: `E |= {group[i] for i in pm[:k]}` печаталось как `E |= group[i] for i in pm[:k]`.
Поэтому строки блока кода не проходят ни demath, ни разбор разметки.

ЗАПУСК:
    python stend/runs_v2/md_to_docx.py ВХОД.md ВЫХОД.docx
"""
import os, re, sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, Mm

FONT = "Times New Roman"
SIZE = Pt(14)


SUP = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")


def demath(s):
    """Снять формульную разметку, сохранив читаемость числа."""
    s = re.sub(r"\$([^$]*)\$", r"\1", s)          # $...$ -> ...
    s = s.replace("{,}", ",").replace("\\cdot", "×")
    # степени: 10^{-8} и 10^-8 -> 10⁻⁸
    s = re.sub(r"\^\{?(-?\d+)\}?", lambda m: m.group(1).translate(SUP), s)
    s = s.replace("^{n}", "\u207f").replace("^n", "\u207f")   # буквенный показатель
    s = re.sub(r"\\Delta_\\text\{([^}]*)\}", r"Δ_\1", s)
    s = re.sub(r"\\Delta", "Δ", s)
    s = re.sub(r"([A-Za-z])_\{?(\d\w*)\}?", r"\1\2", s)   # P_0 -> P0
    s = s.replace("\\", "").replace("^{-", "^-").replace("}", "").replace("{", "")
    return s


def add_runs(p, text):
    """Разложить **жирный** и *курсив* на отдельные фрагменты."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        r = p.add_run(re.sub(r"^(\*\*|\*|`)|(\*\*|\*|`)$", "", part))
        r.font.name, r.font.size = FONT, SIZE
        if part.startswith("**"):
            r.bold = True
        elif part.startswith("*"):
            r.italic = True
        elif part.startswith("`"):
            r.font.name = "Consolas"


def body(doc, text, first_line=True, center=False, italic=False, size=SIZE):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if first_line:
            pf.first_line_indent = Cm(1.25)
    add_runs(p, text)
    for r in p.runs:
        r.font.size = size
        if italic:
            r.italic = True
    return p


def code_block(doc, lines):
    """Листинг: по строке на абзац, дословно, моноширинным шрифтом.

    Ни demath, ни разбор разметки к строкам не применяются: в коде фигурные скобки,
    обратная косая черта и звёздочки — значащие символы, а не оформление.
    """
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(ln.replace("\t", "    "))
        r.font.name = "Consolas"
        r.font.size = Pt(9)


def table(doc, rows):
    head = [c.strip() for c in rows[0].strip("|").split("|")]
    data = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], demath(h))
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(12)
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row[:len(head)]):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], demath(v))
            for r in cells[i].paragraphs[0].runs:
                r.font.size = Pt(12)


def convert(src, dst):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, SIZE
    # Формат страницы и поля по ГОСТ 7.32: A4, левое 30 мм, правое 15 мм, верхнее и
    # нижнее по 20 мм. Шаблон python-docx по умолчанию даёт Letter с дюймовыми полями,
    # и собранный из таких частей документ уезжает с формата.
    for sec in doc.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)

    lines = open(src, encoding="utf-8").read().split("\n")
    i, n_tab, n_head, n_fig, n_code = 0, 0, 0, 0, 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if not s:
            i += 1
            continue

        if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            table(doc, block)
            n_tab += 1
            continue

        if s.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1                       # закрывающий забор
            code_block(doc, block)
            n_code += 1
            continue

        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            # в заголовке снимаются знаки разметки: обратные апострофы и звёздочки
            # иначе они попадают и в сам заголовок, и в оглавление
            txt = re.sub(r"[`*]", "", demath(s.lstrip("# ").strip()))
            h = doc.add_heading(txt, level=min(lvl, 3))
            for r in h.runs:
                r.font.name = FONT
            n_head += 1
            i += 1
            continue

        if s.startswith("[РИСУНОК"):
            body(doc, demath(s), first_line=False, center=True, italic=True, size=Pt(12))
            n_fig += 1
            i += 1
            continue

        if s.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            p = body(doc, demath(" ".join(block)), first_line=False, size=Pt(12))
            p.paragraph_format.left_indent = Cm(1.0)
            for r in p.runs:
                r.italic = True
            continue

        # обычный абзац: склеить до пустой строки
        block = []
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("#", "|", ">", "```", "[РИСУНОК")):
            block.append(lines[i].strip())
            i += 1
        body(doc, demath(" ".join(block)))

    doc.save(dst)
    return n_head, n_tab, n_fig, n_code


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    h, t, f, c = convert(src, dst)
    print(f"{dst}: заголовков {h}, таблиц {t}, мест под рисунки {f}, "
          f"блоков кода {c}, {os.path.getsize(dst)/1024:.0f} КБ")
